from docx import Document
import pandas as pd
import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter import font
import threading
from pathlib import Path

class GeneradorReportesGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("📊 Generador de Reportes Automático")
        self.root.geometry("800x600")
        self.root.configure(bg='#f0f0f0')
        
        # Variables
        self.archivo_csv = tk.StringVar()
        self.template_docx = tk.StringVar()
        self.carpeta_salida = tk.StringVar(value="reportes")
        
        # Configurar estilo
        self.configurar_estilos()
        
        # Crear interfaz
        self.crear_interfaz()
        
        # Centrar ventana
        self.centrar_ventana()
    
    def configurar_estilos(self):
        """Configura los estilos de la aplicación"""
        self.style = ttk.Style()
        self.style.theme_use('clam')
        
        # Configurar colores
        self.style.configure('Title.TLabel', 
                           font=('Arial', 16, 'bold'),
                           background='#f0f0f0',
                           foreground='#2c3e50')
        
        self.style.configure('Subtitle.TLabel',
                           font=('Arial', 10, 'bold'),
                           background='#f0f0f0',
                           foreground='#34495e')
        
        self.style.configure('Custom.TButton',
                           font=('Arial', 10),
                           padding=10)
        
        self.style.configure('Success.TButton',
                           background='#27ae60',
                           foreground='white',
                           font=('Arial', 12, 'bold'))
        
        self.style.map('Success.TButton',
                      background=[('active', '#2ecc71')])
    
    def crear_interfaz(self):
        """Crea toda la interfaz de usuario"""
        # Frame principal
        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # Configurar grid
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        main_frame.columnconfigure(1, weight=1)
        
        # Título principal
        titulo = ttk.Label(main_frame, text="📊 Generador de Reportes", 
                          style='Title.TLabel')
        titulo.grid(row=0, column=0, columnspan=3, pady=(0, 20))
        
        # Sección 1: Archivo CSV
        self.crear_seccion_archivo(main_frame, 1, "📄 Archivo CSV:", 
                                 self.archivo_csv, self.seleccionar_csv,
                                 "Selecciona el archivo CSV con los datos de los clientes")
        
        # Sección 2: Template DOCX
        self.crear_seccion_archivo(main_frame, 3, "📝 Template DOCX:", 
                                 self.template_docx, self.seleccionar_template,
                                 "Selecciona el template de Word (.docx)")
        
        # Sección 3: Carpeta de salida
        self.crear_seccion_carpeta(main_frame, 5)
        
        # Vista previa de datos
        self.crear_vista_previa(main_frame, 7)
        
        # Botones de acción
        self.crear_botones_accion(main_frame, 9)
        
        # Barra de progreso
        self.progress = ttk.Progressbar(main_frame, mode='determinate')
        self.progress.grid(row=10, column=0, columnspan=3, sticky=(tk.W, tk.E), 
                          pady=(10, 0))
        
        # Etiqueta de estado
        self.estado_label = ttk.Label(main_frame, text="Listo para generar reportes")
        self.estado_label.grid(row=11, column=0, columnspan=3, pady=(5, 0))
    
    def crear_seccion_archivo(self, parent, row, titulo, variable, comando, tooltip):
        """Crea una sección para seleccionar archivos"""
        # Etiqueta
        label = ttk.Label(parent, text=titulo, style='Subtitle.TLabel')
        label.grid(row=row, column=0, sticky=tk.W, pady=(10, 5))
        
        # Entry para mostrar ruta
        entry = ttk.Entry(parent, textvariable=variable, state='readonly', width=50)
        entry.grid(row=row+1, column=0, columnspan=2, sticky=(tk.W, tk.E), padx=(0, 10))
        
        # Botón examinar
        btn = ttk.Button(parent, text="📁 Examinar", command=comando,
                        style='Custom.TButton')
        btn.grid(row=row+1, column=2, sticky=tk.W)
        
        # Tooltip
        self.crear_tooltip(entry, tooltip)
    
    def crear_seccion_carpeta(self, parent, row):
        """Crea la sección para la carpeta de salida"""
        label = ttk.Label(parent, text="📂 Carpeta de salida:", style='Subtitle.TLabel')
        label.grid(row=row, column=0, sticky=tk.W, pady=(10, 5))
        
        entry = ttk.Entry(parent, textvariable=self.carpeta_salida, width=50)
        entry.grid(row=row+1, column=0, columnspan=2, sticky=(tk.W, tk.E), padx=(0, 10))
        
        btn = ttk.Button(parent, text="📁 Cambiar", command=self.seleccionar_carpeta,
                        style='Custom.TButton')
        btn.grid(row=row+1, column=2, sticky=tk.W)
        
        self.crear_tooltip(entry, "Carpeta donde se guardarán los reportes generados")
    
    def crear_vista_previa(self, parent, row):
        """Crea la vista previa de datos"""
        label = ttk.Label(parent, text="👁️ Vista previa de datos:", style='Subtitle.TLabel')
        label.grid(row=row, column=0, sticky=tk.W, pady=(20, 5))
        
        # Frame para la tabla
        table_frame = ttk.Frame(parent)
        table_frame.grid(row=row+1, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S),
                        pady=(0, 10))
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        
        # Treeview para mostrar datos
        self.tree = ttk.Treeview(table_frame, height=6)
        self.tree.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # Scrollbars
        v_scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        v_scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        self.tree.configure(yscrollcommand=v_scrollbar.set)
        
        h_scrollbar = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        h_scrollbar.grid(row=1, column=0, sticky=(tk.W, tk.E))
        self.tree.configure(xscrollcommand=h_scrollbar.set)
        
        parent.rowconfigure(row+1, weight=1)
    
    def crear_botones_accion(self, parent, row):
        """Crea los botones de acción"""
        button_frame = ttk.Frame(parent)
        button_frame.grid(row=row, column=0, columnspan=3, pady=20)
        
        # Botón generar reportes
        self.btn_generar = ttk.Button(button_frame, text="🚀 Generar Reportes",
                                     command=self.generar_reportes_thread,
                                     style='Success.TButton')
        self.btn_generar.pack(side=tk.LEFT, padx=(0, 10))
        
        # Botón abrir carpeta
        self.btn_abrir = ttk.Button(button_frame, text="📂 Abrir Carpeta",
                                   command=self.abrir_carpeta_salida,
                                   style='Custom.TButton')
        self.btn_abrir.pack(side=tk.LEFT, padx=(0, 10))
        
        # Botón limpiar
        btn_limpiar = ttk.Button(button_frame, text="🗑️ Limpiar",
                               command=self.limpiar_campos,
                               style='Custom.TButton')
        btn_limpiar.pack(side=tk.LEFT)
    
    def crear_tooltip(self, widget, text):
        """Crea un tooltip para un widget"""
        def on_enter(event):
            tooltip = tk.Toplevel()
            tooltip.wm_overrideredirect(True)
            tooltip.wm_geometry(f"+{event.x_root+10}+{event.y_root+10}")
            label = tk.Label(tooltip, text=text, background="lightyellow",
                           font=("Arial", 9), wraplength=200)
            label.pack()
            widget.tooltip = tooltip
        
        def on_leave(event):
            if hasattr(widget, 'tooltip'):
                widget.tooltip.destroy()
                del widget.tooltip
        
        widget.bind("<Enter>", on_enter)
        widget.bind("<Leave>", on_leave)
    
    def centrar_ventana(self):
        """Centra la ventana en la pantalla"""
        self.root.update_idletasks()
        width = self.root.winfo_width()
        height = self.root.winfo_height()
        x = (self.root.winfo_screenwidth() // 2) - (width // 2)
        y = (self.root.winfo_screenheight() // 2) - (height // 2)
        self.root.geometry(f"{width}x{height}+{x}+{y}")
    
    def seleccionar_csv(self):
        """Selecciona el archivo CSV"""
        archivo = filedialog.askopenfilename(
            title="Seleccionar archivo CSV",
            filetypes=[("Archivos CSV", "*.csv"), ("Todos los archivos", "*.*")]
        )
        if archivo:
            self.archivo_csv.set(archivo)
            self.cargar_vista_previa()
    
    def seleccionar_template(self):
        """Selecciona el template DOCX"""
        archivo = filedialog.askopenfilename(
            title="Seleccionar template DOCX",
            filetypes=[("Archivos Word", "*.docx"), ("Todos los archivos", "*.*")]
        )
        if archivo:
            self.template_docx.set(archivo)
    
    def seleccionar_carpeta(self):
        """Selecciona la carpeta de salida"""
        carpeta = filedialog.askdirectory(title="Seleccionar carpeta de salida")
        if carpeta:
            self.carpeta_salida.set(carpeta)
    
    def cargar_vista_previa(self):
        """Carga la vista previa de los datos CSV"""
        try:
            if not self.archivo_csv.get():
                return
            
            # Leer CSV
            df = pd.read_csv(self.archivo_csv.get())
            
            # Limpiar tabla
            for item in self.tree.get_children():
                self.tree.delete(item)
            
            # Configurar columnas
            self.tree["columns"] = list(df.columns)
            self.tree["show"] = "headings"
            
            # Configurar encabezados
            for col in df.columns:
                self.tree.heading(col, text=col)
                self.tree.column(col, width=100, minwidth=50)
            
            # Insertar datos (máximo 10 filas para vista previa)
            for index, row in df.head(10).iterrows():
                self.tree.insert("", tk.END, values=list(row))
            
            # Mostrar información
            total_registros = len(df)
            self.estado_label.config(text=f"CSV cargado: {total_registros} registros encontrados")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error al cargar CSV: {str(e)}")
    
    def reemplazar_placeholder(self, paragraph, variables):
        """Reemplaza placeholders en un párrafo"""
        runs = paragraph.runs
        i = 0
        while i < len(runs):
            # Intentar juntar hasta 3 runs consecutivos
            texto_junto = runs[i].text
            if i + 1 < len(runs):
                texto_junto += runs[i + 1].text
            if i + 2 < len(runs):
                texto_junto += runs[i + 2].text
            
            for key, value in variables.items():
                if key in texto_junto:
                    # Reemplazar el placeholder
                    remaining = texto_junto.replace(key, value)
                    runs[i].text = remaining
                    if i + 1 < len(runs):
                        runs[i + 1].text = ""
                    if i + 2 < len(runs):
                        runs[i + 2].text = ""
                    break
            i += 1
    
    def generar_reportes_thread(self):
        """Ejecuta la generación de reportes en un hilo separado"""
        thread = threading.Thread(target=self.generar_reportes)
        thread.daemon = True
        thread.start()
    
    def generar_reportes(self):
        """Genera los reportes automáticamente"""
        try:
            # Validar campos
            if not self.archivo_csv.get():
                messagebox.showerror("Error", "Selecciona un archivo CSV")
                return
            
            if not self.template_docx.get():
                messagebox.showerror("Error", "Selecciona un template DOCX")
                return
            
            if not os.path.exists(self.archivo_csv.get()):
                messagebox.showerror("Error", "El archivo CSV no existe")
                return
            
            if not os.path.exists(self.template_docx.get()):
                messagebox.showerror("Error", "El template DOCX no existe")
                return
            
            # Actualizar estado
            self.estado_label.config(text="Iniciando generación de reportes...")
            self.btn_generar.config(state='disabled')
            
            # Leer datos
            df = pd.read_csv(self.archivo_csv.get())
            total_clientes = len(df)
            
            # Crear carpeta de salida
            carpeta_salida = self.carpeta_salida.get()
            if not os.path.exists(carpeta_salida):
                os.makedirs(carpeta_salida)
            
            # Configurar barra de progreso
            self.progress.config(maximum=total_clientes, value=0)
            
            reportes_generados = 0
            
            # Generar reportes
            for index, cliente in df.iterrows():
                try:
                    # Actualizar estado
                    self.estado_label.config(text=f"Generando reporte {index + 1} de {total_clientes}...")
                    
                    # Cargar template
                    doc = Document(self.template_docx.get())
                    
                    # Crear variables de reemplazo
                    variables = {f"${col}$": str(cliente[col]) for col in df.columns}
                    
                    # Reemplazar placeholders
                    for p in doc.paragraphs:
                        self.reemplazar_placeholder(p, variables)
                    
                    # Generar nombre de archivo
                    primer_campo = str(cliente.iloc[0]).replace(' ', '_')
                    primer_campo = "".join(c for c in primer_campo if c.isalnum() or c in ('_', '-'))
                    nombre_archivo = f"{carpeta_salida}/reporte_{primer_campo}.docx"
                    
                    # Guardar documento
                    doc.save(nombre_archivo)
                    reportes_generados += 1
                    
                    # Actualizar barra de progreso
                    self.progress.config(value=index + 1)
                    self.root.update_idletasks()
                    
                except Exception as e:
                    print(f"Error generando reporte para cliente {index}: {str(e)}")
                    continue
            
            # Finalizar
            self.estado_label.config(text=f"✅ ¡Completado! {reportes_generados} reportes generados")
            messagebox.showinfo("Éxito", 
                              f"Se generaron {reportes_generados} reportes exitosamente\n"
                              f"Ubicación: {carpeta_salida}")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error durante la generación: {str(e)}")
            self.estado_label.config(text="❌ Error en la generación")
        
        finally:
            self.btn_generar.config(state='normal')
            self.progress.config(value=0)
    
    def abrir_carpeta_salida(self):
        """Abre la carpeta de salida en el explorador"""
        carpeta = self.carpeta_salida.get()
        if os.path.exists(carpeta):
            if os.name == 'nt':  # Windows
                os.startfile(carpeta)
            elif os.name == 'posix':  # macOS y Linux
                os.system(f'open "{carpeta}"' if sys.platform == 'darwin' else f'xdg-open "{carpeta}"')
        else:
            messagebox.showwarning("Advertencia", "La carpeta no existe aún")
    
    def limpiar_campos(self):
        """Limpia todos los campos"""
        self.archivo_csv.set("")
        self.template_docx.set("")
        self.carpeta_salida.set("reportes")
        
        # Limpiar tabla
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        self.estado_label.config(text="Campos limpiados - Listo para empezar")
        self.progress.config(value=0)

def main():
    root = tk.Tk()
    app = GeneradorReportesGUI(root)
    root.mainloop()

if __name__ == "__main__":
    main()