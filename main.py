from docx import Document
import pandas as pd
import os

# Leer datos desde CSV
df = pd.read_csv('clientes.csv')

# Crear carpeta para reportes si no existe
if not os.path.exists('reportes'):
    os.mkdir('reportes')

def reemplazar_placeholder(paragraph, variables):
    runs = paragraph.runs
    i = 0
    while i < len(runs):
        # Intentar juntar hasta 3 runs consecutivos (ajusta si tus placeholders son más largos)
        texto_junto = runs[i].text
        if i + 1 < len(runs):
            texto_junto += runs[i + 1].text
        if i + 2 < len(runs):
            texto_junto += runs[i + 2].text
        
        for key, value in variables.items():
            if key in texto_junto:
                # Reemplazar el placeholder en los runs correspondientes
                remaining = texto_junto.replace(key, value)
                # Vaciar los runs usados
                runs[i].text = remaining
                if i + 1 < len(runs):
                    runs[i + 1].text = ""
                if i + 2 < len(runs):
                    runs[i + 2].text = ""
                break  # Una vez reemplazado, pasamos al siguiente run
        i += 1

# Generar reportes para cada cliente
for index, cliente in df.iterrows():
    doc = Document('template.docx')
    
    variables = {f"${col}$": str(cliente[col]) for col in df.columns}
    
    for p in doc.paragraphs:
        reemplazar_placeholder(p, variables)
    
    # Guardar el reporte con nombre del cliente
    nombre_archivo = f"reportes/reporte_{cliente.iloc[0].replace(' ', '_')}.docx"
    doc.save(nombre_archivo)
    print(f"Reporte generado: {nombre_archivo}")